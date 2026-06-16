import os
import re
import time
import base64
import threading
import configparser
import gc

from PySide6.QtWidgets import (
    QApplication, QWidget, QLabel, QPushButton,
    QLineEdit, QFileDialog, QCheckBox, QTextEdit, QProgressBar,
    QVBoxLayout, QHBoxLayout, QGroupBox, QMessageBox, QComboBox,
    QSpinBox
)
from PySide6.QtCore import Signal, QObject, QThread

from pitcsuite.ui_helpers import lbl, hline, START_BTN, STOP_BTN
from pitcsuite.templates import download_template
from pitcsuite.config import config_path as _config_path

class BillWorkerSignals(QObject):
    log      = Signal(str)
    progress = Signal(int, int)
    finished = Signal()


class BillSoftPanel(QWidget):
    def __init__(self):
        super().__init__()
        self.stop_flag = False; self.pause_flag = False
        self.signals = BillWorkerSignals()
        self.signals.log.connect(lambda m: self.log_box.append(m))
        self.signals.progress.connect(self.update_progress)
        self.signals.finished.connect(self.on_finished)

        layout = QVBoxLayout(self)
        layout.setSpacing(10)
        layout.addWidget(lbl("E/Bills Downloader", bold=True, color="#7eb8f7"))
        layout.addWidget(hline())

        g = QGroupBox("Inputs")
        iv = QVBoxLayout(g)
        iv.addLayout(self._excel_row())
        iv.addLayout(self._url_row())
        iv.addLayout(self._out_row())
        layout.addWidget(g)

        og = QGroupBox("Options")
        oh = QHBoxLayout(og)
        self.chk_sr       = QCheckBox("Add Sr. No. on Bill")
        self.chk_sep      = QCheckBox("Separate PDFs")
        self.chk_headless = QCheckBox("Headless (background)"); self.chk_headless.setChecked(True)
        oh.addWidget(self.chk_sr); oh.addWidget(self.chk_sep); oh.addWidget(self.chk_headless); oh.addStretch()
        layout.addWidget(og)

        bh = QHBoxLayout()
        self.btn_start = QPushButton("▶  START"); self.btn_start.setStyleSheet(START_BTN)
        self.btn_pause = QPushButton("⏸  PAUSE"); self.btn_pause.setStyleSheet("background:#a80; border:none; border-radius:4px; padding:6px 14px; color:#fff;")
        self.btn_start.clicked.connect(self.start_stop); self.btn_pause.clicked.connect(self.pause_resume)
        bh.addWidget(self.btn_start); bh.addWidget(self.btn_pause); bh.addStretch()
        layout.addLayout(bh)

        ph = QHBoxLayout()
        self.lbl_prog = QLabel("0 / 0 bills"); ph.addWidget(self.lbl_prog); ph.addStretch()
        layout.addLayout(ph)
        self.bar = QProgressBar(); layout.addWidget(self.bar)

        lg = QGroupBox("Logs")
        lv = QVBoxLayout(lg)
        self.log_box = QTextEdit(); self.log_box.setReadOnly(True); self.log_box.setFixedHeight(100)
        lv.addWidget(self.log_box)
        layout.addWidget(lg)

    def _excel_row(self):
        r = QHBoxLayout(); r.addWidget(QLabel("Excel File:"))
        self.ed_excel = QLineEdit(); self.ed_excel.setReadOnly(True); self.ed_excel.setProperty("preferred_width", 280); r.addWidget(self.ed_excel, 1)
        b = QPushButton("Browse"); b.clicked.connect(lambda: (p := QFileDialog.getOpenFileName(self,"Excel","","Excel (*.xlsx)")[0]) and self.ed_excel.setText(p)); r.addWidget(b)
        return r

    def _url_row(self):
        r = QHBoxLayout(); r.addWidget(QLabel("Bills URL:"))
        self.ed_url = QLineEdit("https://bill.pitc.com.pk/pescobill/"); self.ed_url.setReadOnly(True); self.ed_url.setProperty("preferred_width", 310); r.addWidget(self.ed_url, 1)
        self.btn_url = QPushButton("Edit"); self.btn_url.clicked.connect(self.toggle_url); r.addWidget(self.btn_url)
        return r

    def _out_row(self):
        r = QHBoxLayout(); r.addWidget(QLabel("Output Folder:"))
        self.ed_out = QLineEdit("D:\\BillsPDF"); self.ed_out.setReadOnly(True); r.addWidget(self.ed_out, 1)
        o = QPushButton("Open");  o.clicked.connect(lambda: os.startfile(self.ed_out.text())); r.addWidget(o)
        c = QPushButton("Change"); c.clicked.connect(lambda: (p := QFileDialog.getExistingDirectory(self,"Output")) and self.ed_out.setText(p)); r.addWidget(c)
        return r

    def toggle_url(self):
        if self.btn_url.text() == "Edit":
            self.ed_url.setReadOnly(False); self.btn_url.setText("Save")
        else:
            self.ed_url.setReadOnly(True); self.btn_url.setText("Edit")

    def start_stop(self):
        if self.btn_start.text().endswith("START"):
            if not self.ed_excel.text():
                QMessageBox.warning(self, "Error", "Select Excel file first"); return
            self.stop_flag = False; self.pause_flag = False
            self.btn_start.setText("■  STOP"); self.btn_start.setStyleSheet(STOP_BTN)
            threading.Thread(target=self.run_worker, daemon=True).start()
        else:
            self.stop_flag = True; self.btn_start.setText("▶  START"); self.btn_start.setStyleSheet(START_BTN)

    def pause_resume(self):
        self.pause_flag = not self.pause_flag
        self.btn_pause.setText("▶  RESUME" if self.pause_flag else "⏸  PAUSE")

    def update_progress(self, cur, total):
        self.bar.setMaximum(total); self.bar.setValue(cur)
        self.lbl_prog.setText(f"{cur} / {total} bills")

    def on_finished(self):
        self.btn_start.setText("▶  START"); self.btn_start.setStyleSheet(START_BTN)
        self.btn_pause.setText("⏸  PAUSE")

        QMessageBox.information(self, "Done", "E/Bills Downloader completed.")

    def _unique_pdf(self, folder, base):
        p = os.path.join(folder, f"{base}.pdf")
        if not os.path.exists(p): return p
        i = 1
        while True:
            p2 = os.path.join(folder, f"{base} ({i}).pdf"); 
            if not os.path.exists(p2): return p2
            i += 1

    def _inject_sr(self, driver, sr):
        driver.execute_script(f"""
        let d=document.getElementById('sr-overlay'); if(d) d.remove();
        d=document.createElement('div'); d.id='sr-overlay'; d.innerText='Sr. No. {sr}';
        d.style.cssText='position:fixed;top:10px;right:20px;font-size:20px;font-weight:bold;text-decoration:underline;z-index:99999;';
        document.body.appendChild(d);""")

    def _fit_pdf_to_a4(self, pdf_path, sr=None):
        import fitz

        src = fitz.open(pdf_path)
        if not src.page_count:
            src.close()
            return

        src_page = src[0]
        content_rect = None

        for block in src_page.get_text("blocks"):
            if len(block) >= 5 and str(block[4]).strip():
                rect = fitz.Rect(block[:4])
                content_rect = rect if content_rect is None else content_rect | rect

        if content_rect is None or content_rect.is_empty:
            content_rect = src_page.rect
        else:
            content_rect = content_rect + (-8, -8, 8, 8)
            content_rect &= src_page.rect

        out = fitz.open()
        a4 = fitz.paper_rect("a4")
        page = out.new_page(width=a4.width, height=a4.height)

        if sr:
            stamp = f"Sr. No. {sr}"
            text_width = fitz.get_text_length(stamp, fontname="helv", fontsize=14)
            page.insert_text(
                fitz.Point(a4.width - text_width - 24, 24),
                stamp,
                fontsize=14,
                fontname="helv",
                color=(0, 0, 0),
            )
            target = fitz.Rect(18, 36, a4.width - 18, a4.height - 18)
        else:
            target = fitz.Rect(18, 18, a4.width - 18, a4.height - 18)

        page.show_pdf_page(target, src, 0, clip=content_rect, keep_proportion=False)

        base, ext = os.path.splitext(pdf_path)
        tmp_path = f"{base}_tmp{ext}"
        out.save(tmp_path, garbage=4, deflate=True)
        out.close()
        src.close()
        del out, src
        gc.collect()

        for attempt in range(5):
            try:
                os.replace(tmp_path, pdf_path)
                break
            except PermissionError:
                if attempt == 4:
                    raise
                time.sleep(0.2)

    def run_worker(self):
        try:
            from selenium import webdriver
            from selenium.webdriver.common.by import By
            from selenium.webdriver.chrome.service import Service
            from selenium.webdriver.support.ui import WebDriverWait
            from selenium.webdriver.support import expected_conditions as EC
            from selenium.common.exceptions import TimeoutException
            from webdriver_manager.chrome import ChromeDriverManager
            from openpyxl import load_workbook
            from pypdf import PdfReader, PdfWriter
        except ImportError as e:
            self.signals.log.emit(f"ERROR: Missing library — {e}"); self.signals.finished.emit(); return

        excel = self.ed_excel.text(); out = self.ed_out.text(); url = self.ed_url.text()
        os.makedirs(out, exist_ok=True)
        wb = load_workbook(excel); sh = wb.active

        opts = webdriver.ChromeOptions()
        if self.chk_headless.isChecked(): opts.add_argument("--headless=new")
        driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=opts)

        total = sh.max_row - 1; done = 0; collected = []
        WAIT = 12; MAX_ATTEMPTS = 3; MIN_RENDER_WAIT = 5

        def visible_search_boxes():
            return [el for el in driver.find_elements(By.ID, "searchTextBox") if el.is_displayed()]

        def page_text():
            try:
                return driver.find_element(By.TAG_NAME, "body").text
            except Exception:
                return ""

        def bill_state(ac_no):
            text = page_text()
            text_l = text.lower()
            if "bill not found" in text_l:
                return "not_found"
            if "does not belongs" in text_l:
                return "wrong_ac"

            has_old_print_button = bool(driver.find_elements(By.CSS_SELECTOR, ".auto-center.btn.btn-secondary"))
            has_print_command = bool(driver.find_elements(
                By.XPATH,
                "//*[self::button or self::a or self::input][contains("
                "translate(normalize-space(.), 'ABCDEFGHIJKLMNOPQRSTUVWXYZ', 'abcdefghijklmnopqrstuvwxyz'), 'print') "
                "or contains(translate(@value, 'ABCDEFGHIJKLMNOPQRSTUVWXYZ', 'abcdefghijklmnopqrstuvwxyz'), 'print') "
                "or contains(translate(@onclick, 'ABCDEFGHIJKLMNOPQRSTUVWXYZ', 'abcdefghijklmnopqrstuvwxyz'), 'print')]"
            ))

            bill_keywords = [
                "reference", "consumer", "billing month", "due date",
                "payable", "amount", "pesco"
            ]
            keyword_hits = sum(1 for keyword in bill_keywords if keyword in text_l)
            ac_visible = ac_no in re.sub(r"\s+", "", text)

            if has_old_print_button or has_print_command:
                return "loaded"
            if ac_visible and keyword_hits >= 2:
                return "loaded"
            if not visible_search_boxes() and keyword_hits >= 2:
                return "loaded"
            return "loading"

        def wait_until_render_ready(clicked_at):
            remaining = MIN_RENDER_WAIT - (time.time() - clicked_at)
            if remaining > 0:
                time.sleep(remaining)

            try:
                WebDriverWait(driver, 5).until(lambda d: d.execute_script("""
                    const ready = document.readyState === 'complete';
                    const imagesReady = Array.from(document.images || []).every(img => img.complete);
                    return ready && imagesReady;
                """))
            except Exception:
                pass

        for r in range(2, sh.max_row + 1):
            if self.stop_flag:
                sh.cell(r, 3).value = "Stopped"; wb.save(excel); break
            while self.pause_flag: time.sleep(0.3)
            ac = sh.cell(r, 1).value; sr = sh.cell(r, 2).value
            if not ac: break
            ac = str(ac).strip(); self.signals.log.emit(f"Processing AC: {ac}")
            status = None; bill_loaded = False

            for attempt in range(1, MAX_ATTEMPTS + 1):
                driver.get(url); time.sleep(2)
                try:
                    wait = WebDriverWait(driver, 20)
                    box = wait.until(EC.element_to_be_clickable((By.ID, "searchTextBox")))
                    box.clear(); box.send_keys(ac)
                    clicked_at = time.time()
                    driver.find_element(By.ID, "btnSearch").click()

                    end = clicked_at + WAIT
                    state = "loading"
                    while time.time() < end:
                        state = bill_state(ac)
                        if state == "loaded" and time.time() - clicked_at < MIN_RENDER_WAIT:
                            state = "loading"
                        if state != "loading":
                            break
                        time.sleep(0.5)
                except TimeoutException:
                    state = "loading"

                if state == "not_found":
                    status = "Bill not Found"; break
                if state == "wrong_ac":
                    status = "Wrong AC No."; break
                if state == "loaded":
                    wait_until_render_ready(clicked_at)
                    bill_loaded = True; status = "Saved"; break
                self.signals.log.emit(f"Attempt {attempt}/{MAX_ATTEMPTS}: bill page still loading")

            if status is None: status = "Bill not loaded (timeout/site changed)"

            if not bill_loaded:
                sh.cell(r, 3).value = status; wb.save(excel); done += 1
                self.signals.progress.emit(done, total); self.signals.log.emit(f"→ {status}"); continue

            pdf = driver.execute_cdp_cmd("Page.printToPDF", {
                "printBackground": True, "paperWidth": 8.27, "paperHeight": 11.69,
                "marginTop": 0.15, "marginBottom": 0.15, "marginLeft": 0.15, "marginRight": 0.15, "scale": 0.60
            })
            pdf_name = str(sr).strip() if sr else ac
            p = self._unique_pdf(out, pdf_name)
            with open(p, "wb") as f: f.write(base64.b64decode(pdf["data"]))
            self._fit_pdf_to_a4(p, sr if self.chk_sr.isChecked() and sr else None)
            collected.append((int(sr) if sr and str(sr).isdigit() else 999999, p))
            sh.cell(r, 3).value = "Saved"; self.signals.log.emit("→ PDF created"); wb.save(excel)
            done += 1; self.signals.progress.emit(done, total)

        driver.quit()
        if not self.chk_sep.isChecked() and collected:
            if self.chk_sr.isChecked(): collected.sort(key=lambda x: x[0])
            merged = self._unique_pdf(out, os.path.splitext(os.path.basename(excel))[0])
            from pypdf import PdfWriter, PdfReader
            writer = PdfWriter()
            for _, pp in collected:
                for pg in PdfReader(pp).pages: writer.add_page(pg)
            with open(merged, "wb") as f: writer.write(f)
            for _, pp in collected: os.remove(pp)

        self.signals.finished.emit()


# ──────────────────────────────────────────────────────────────
#  TOOL 3 — Sub-Division Extractor (extractSd.py)
# ──────────────────────────────────────────────────────────────
