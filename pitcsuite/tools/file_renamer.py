import shutil
import uuid
from pathlib import Path

from PySide6.QtCore import QThread, Signal
from PySide6.QtWidgets import (
    QFileDialog, QGroupBox, QHBoxLayout, QLabel, QLineEdit, QMessageBox,
    QPushButton, QTextEdit, QVBoxLayout, QWidget,
)

from pitcsuite.ui_helpers import START_BTN, hline, lbl, notify


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _target_name(original, new_name):
    """Keep the source extension for the renamed file."""
    new_path = Path(new_name)
    source_suffix = Path(original).suffix
    return new_path.with_suffix(source_suffix).name if source_suffix else new_name


class FileRenamerWorker(QThread):
    progress = Signal(str)
    done = Signal(int, int, list)
    failed = Signal(str)

    def __init__(self, excel_path, source_dir, destination_dir=""):
        super().__init__()
        self.excel_path = Path(excel_path)
        self.source_dir = Path(source_dir)
        self.destination_dir = Path(destination_dir) if destination_dir else None

    def run(self):
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(self.excel_path)
            sheet = workbook.active
            mappings = []
            for row in range(2, sheet.max_row + 1):
                original = _cell_text(sheet.cell(row, 1).value)
                new_name = _cell_text(sheet.cell(row, 2).value)
                if not original and not new_name:
                    continue
                if not original or not new_name:
                    raise ValueError(f"Row {row}: both columns A and B are required.")
                if Path(new_name).name != new_name:
                    raise ValueError(f"Row {row}: new name must be a file name, not a path.")
                mappings.append((original, new_name, row))
            if not mappings:
                workbook.close()
                raise ValueError("No mappings found in columns A and B from row 2.")

            files = {path.name.casefold(): path for path in self.source_dir.iterdir() if path.is_file()}
            stems = {}
            for path in files.values():
                stems.setdefault(path.stem.casefold(), []).append(path)

            resolved, issues, remarks = [], [], {}
            for original, requested_name, row in mappings:
                source = files.get(original.casefold())
                if source is None:
                    candidates = stems.get(Path(original).stem.casefold(), [])
                    source = candidates[0] if len(candidates) == 1 else None
                if source is None:
                    issues.append(f"Row {row}: source not found - {original}")
                    remarks[row] = "Not Found"
                else:
                    # Use the actual source file's suffix, including when
                    # column A contains only a filename stem.
                    resolved.append((source, _target_name(source.name, requested_name), row))
            if not resolved:
                for row, remark in remarks.items():
                    sheet.cell(row, 4).value = remark
                workbook.save(self.excel_path)
                workbook.close()
                self.done.emit(0, len(issues), issues)
                return

            target_dir = self.destination_dir or self.source_dir
            target_dir.mkdir(parents=True, exist_ok=True)
            seen = set()
            valid = []
            for source, new_name, row in resolved:
                target = target_dir / new_name
                key = str(target).casefold()
                if key in seen:
                    issues.append(f"Row {row}: duplicate target name - {new_name}")
                    remarks[row] = "Duplicate Target"
                else:
                    seen.add(key); valid.append((source, target, row))

            if self.destination_dir:
                completed = self._copy(valid, issues, remarks)
            else:
                completed = self._rename_in_place(valid, issues, remarks)
            for row, remark in remarks.items():
                sheet.cell(row, 4).value = remark
            workbook.save(self.excel_path)
            workbook.close()
            self.done.emit(completed, len(issues), issues)
        except Exception as exc:
            self.failed.emit(str(exc))

    def _copy(self, items, issues, remarks):
        completed = 0
        for index, (source, target, row) in enumerate(items, 1):
            if self.isInterruptionRequested():
                break
            if target.exists():
                issues.append(f"Row {row}: target already exists - {target.name}")
                remarks[row] = "Already Exists"
                continue
            shutil.copy2(source, target)
            completed += 1
            remarks[row] = "Done"
            self.progress.emit(f"Copied {index} of {len(items)}: {source.name} -> {target.name}")
        return completed

    def _rename_in_place(self, items, issues, remarks):
        source_keys = {str(source).casefold() for source, _, _ in items}
        planned = []
        for source, target, row in items:
            if target.exists() and str(target).casefold() not in source_keys:
                issues.append(f"Row {row}: target already exists - {target.name}")
                remarks[row] = "Already Exists"
            elif str(source).casefold() != str(target).casefold():
                planned.append((source, target, row))

        staged = []
        try:
            for source, target, row in planned:
                temporary = source.with_name(f".__pitcsuite_rename_{uuid.uuid4().hex}{source.suffix}")
                source.rename(temporary); staged.append((temporary, target, row))
            for index, (temporary, target, row) in enumerate(staged, 1):
                temporary.rename(target)
                remarks[row] = "Done"
                self.progress.emit(f"Renamed {index} of {len(staged)}: {target.name}")
            for source, target, row in items:
                if source == target:
                    remarks[row] = "Done (Already Named)"
            return len(staged) + sum(1 for source, target, _ in items if source == target)
        except Exception:
            for temporary, target, _ in reversed(staged):
                if temporary.exists() and not target.exists():
                    temporary.rename(target)
            raise


class FileRenamerPanel(QWidget):
    def __init__(self):
        super().__init__(); self.worker = None
        layout = QVBoxLayout(self); layout.setSpacing(10)
        layout.addWidget(lbl("File Renamer", bold=True, color="#7eb8f7"))
        layout.addWidget(lbl("Use Excel column A for original names and column B for new names, starting at row 2. The original extension is kept automatically (PDF stays PDF, XLSX stays XLSX). Remarks are written to column D.", color="#556"))
        layout.addWidget(hline())
        group = QGroupBox("Inputs"); form = QVBoxLayout(group)
        self.excel_ed, self.source_ed, self.destination_ed = QLineEdit(), QLineEdit(), QLineEdit()
        for label, edit, kind in (("Excel file:", self.excel_ed, "excel"), ("Source folder:", self.source_ed, "source"), ("Destination (optional):", self.destination_ed, "destination")):
            row = QHBoxLayout(); row.addWidget(QLabel(label)); edit.setReadOnly(True); edit.setProperty("preferred_width", 245); row.addWidget(edit, 1)
            browse = QPushButton("Browse"); browse.clicked.connect(lambda _, e=edit, k=kind: self._browse(e, k)); row.addWidget(browse); form.addLayout(row)
        layout.addWidget(group)
        self.start_btn = QPushButton("▶  RENAME FILES"); self.start_btn.setStyleSheet(START_BTN); self.start_btn.clicked.connect(self._start); layout.addWidget(self.start_btn)
        self.status = QLabel("Status: Idle"); self.status.setStyleSheet("color:#4a90d9;"); layout.addWidget(self.status)
        self.log_box = QTextEdit(); self.log_box.setReadOnly(True); layout.addWidget(self.log_box)

    def _browse(self, edit, kind):
        path = QFileDialog.getOpenFileName(self, "Excel input file", "", "Excel files (*.xlsx *.xlsm);;All files (*.*)")[0] if kind == "excel" else QFileDialog.getExistingDirectory(self, "Select folder")
        if path: edit.setText(path)

    def _start(self):
        excel, source = Path(self.excel_ed.text().strip()), Path(self.source_ed.text().strip())
        destination = self.destination_ed.text().strip()
        if not excel.is_file() or not source.is_dir():
            QMessageBox.warning(self, "Missing input", "Select an existing Excel file and source folder first."); return
        if destination and Path(destination).resolve() == source.resolve():
            destination = ""
        self.log_box.clear(); self.start_btn.setEnabled(False); self.status.setText("Status: Working...")
        self.worker = FileRenamerWorker(excel, source, destination)
        self.worker.progress.connect(self._progress); self.worker.done.connect(self._done); self.worker.failed.connect(self._failed)
        self.worker.finished.connect(lambda: self.start_btn.setEnabled(True)); self.worker.start()

    def _progress(self, message): self.log_box.append(message)

    def _done(self, completed, issue_count, issues):
        self.status.setText(f"Status: Completed — {completed} file(s) processed")
        message = f"Processed {completed} file(s)." + (("\n\nSkipped / issues:\n" + "\n".join(issues)) if issues else "")
        notify(self, "File Renamer", message)

    def _failed(self, message):
        self.status.setText("Status: Failed"); notify(self, "File Renamer error", message, critical=True)
