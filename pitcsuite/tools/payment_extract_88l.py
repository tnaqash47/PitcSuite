import os
import re
from pathlib import Path

from PySide6.QtCore import QThread, Signal
from PySide6.QtWidgets import (
    QFileDialog, QGroupBox, QHBoxLayout, QLabel, QLineEdit, QMessageBox,
    QProgressBar, QPushButton, QTextEdit, QVBoxLayout, QWidget,
)

from pitcsuite.ui_helpers import hline, lbl, START_BTN, STOP_BTN, notify
from pitcsuite.templates import download_template


def batch_key(filename):
    upper = filename.upper()
    kind = "GEN" if "88L" in upper and "GEN" in upper else "MDI" if "88L" in upper and "MDI" in upper else None
    if not kind:
        return None
    explicit = list(re.finditer(r"(?:^|[^A-Z])B[- _]?(\d{1,2})(?=[^0-9]|$)", filename, re.I))
    match = explicit[-1] if explicit else re.search(r"(?:_|-)(\d{1,2})\.txt$", filename, re.I)
    return f"{kind}-B{int(match.group(1)):02d}" if match else None


def month_year(path, root):
    try:
        parts = path.relative_to(root).parts
    except ValueError:
        parts = path.parts
    for part in parts:
        if re.fullmatch(r"\d{4}-\d{2}", part):
            return part[:4]
    return ""


def discover_files(root_text):
    root = Path(root_text)
    records = []
    for path in root.rglob("*.txt"):
        if path.is_file() and batch_key(path.name):
            records.append({"path": path, "key": batch_key(path.name), "year": month_year(path, root)})
    return records


def ledger_header(text):
    for line in text.splitlines()[:12]:
        if re.search(r"(?:BATCH\s*[: ]|LEDGER|PROCESSING)", line, re.I):
            return line
    return ""


def ledger_year(text, fallback=""):
    header = "\n".join(text.splitlines()[:12])
    match = re.search(r"(?:BILLING\s+MONTH|MONTH(?:\s+OF)?)\b[^\r\n]*?\b((?:19|20)\d{2})\b", header, re.I)
    return match.group(1) if match else fallback


def split_pages(text, header):
    lines = text.splitlines()
    starts = [i for i, line in enumerate(lines) if re.search(r"(?:S/DIV\s*:|S/Div\s*:|Sub\s+Division\b)", line, re.I)]
    return ["\n".join(([header] if header else []) + lines[start:end]) for start, end in zip(starts, starts[1:] + [len(lines)])]


def consumer_reference(ref_no, sub_div=""):
    """Resolve the 7-digit ledger AC from a plain or composite reference."""
    reference = safe_text(ref_no)
    digits = re.sub(r"\D", "", reference)
    if not digits or reference.lower() in {"nan", "none", "null"}:
        return ""
    if len(digits) <= 7:
        return digits.zfill(7)

    # Composite references contain the subdivision followed by the printed
    # account suffix.  Only use this fallback when the selected subdivision is
    # actually present in the reference; never guess from an arbitrary suffix.
    subdivision = re.sub(r"\D", "", safe_text(sub_div))
    if subdivision and subdivision in digits:
        suffix = digits[digits.find(subdivision) + len(subdivision):]
        if suffix:
            return suffix[-6:].zfill(7)
    return ""


def extract_payment(page_text, ref_no, sub_div=""):
    """Return a payment only when *ref_no* is an exact consumer row on page.

    88L ledgers wrap a consumer row over several lines.  The account number is
    the first token of the row and payment amount/date may be on a continuation
    line.  Matching arbitrary text in the page (or allowing an empty reference
    to match) causes unrelated months to be exported.
    """
    reference = consumer_reference(ref_no, sub_div)
    if not reference:
        return None, None

    lines = page_text.splitlines()
    # Excel may supply the AC as a number and drop leading zeroes.  Compare
    # digit-only values after removing leading zeroes, but never compare a
    # partial account number or a substring of another token.
    reference_key = reference.lstrip("0") or "0"

    consumer_row = re.compile(r"^\s*(\d{7})(?=\s|$)")
    payment_pattern = re.compile(r"\b(\d+(?:\.\d+)?)\s+(\d{2}/\d{2})\b")
    for index, line in enumerate(lines):
        row_match = consumer_row.match(line)
        if not row_match:
            continue
        row_key = row_match.group(1).lstrip("0") or "0"
        if row_key != reference_key:
            continue

        end = index + 1
        while end < len(lines) and not consumer_row.match(lines[end]):
            end += 1
        match = payment_pattern.search(" ".join(lines[index:end]))
        if match and float(match.group(1).replace(",", "")) > 0:
            return match.group(1), match.group(2)
    return None, None


def safe_text(value):
    return "" if value is None else str(value).strip()


class PaymentExtract88LWorker(QThread):
    log = Signal(str)
    progress = Signal(int)
    finished = Signal()

    def __init__(self, excel_path, text_root, output_folder):
        super().__init__()
        self.excel_path, self.text_root, self.output_folder = excel_path, text_root, output_folder
        self.stop_flag = False

    def stop(self):
        self.stop_flag = True

    def run(self):
        try:
            from openpyxl import load_workbook
            from PyPDF2 import PdfMerger
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.pdfgen import canvas
        except ImportError as exc:
            self.log.emit(f"ERROR: Missing library — {exc}"); self.finished.emit(); return
        try:
            self.log.emit(f"88L extractor source: {Path(__file__).resolve()}")
            records = discover_files(self.text_root)
            by_batch = {}
            for record in records:
                by_batch.setdefault(record["key"].rsplit("-B", 1)[-1], []).append(record)
            self.log.emit(f"📂 Discovered {len(records)} matching 88L text files recursively.")
            if not records:
                self.log.emit("❌ No supported 88L text files were found."); self.finished.emit(); return

            workbook = load_workbook(self.excel_path)
            sheet = workbook.active
            headers = {safe_text(c.value): c.column for c in sheet[1] if c.value is not None}
            required = ["BN", "Sdiv", "AC No.", "Para Sr. No."]
            missing = [name for name in required if name not in headers]
            if missing:
                self.log.emit("ERROR: Excel is missing columns: " + ", ".join(missing)); self.finished.emit(); return
            status_col = headers.setdefault("Status", sheet.max_column + 1)
            sheet.cell(1, status_col).value = "Status"
            total_col = headers.get("Total Amount")
            if not total_col:
                total_col = headers.get("Found in Files", sheet.max_column + 1)
                sheet.cell(1, total_col).value = "Total Amount"
                headers["Total Amount"] = total_col
            payment_cols = {safe_text(sheet.cell(1, c).value): c for c in range(1, sheet.max_column + 1) if safe_text(sheet.cell(1, c).value).startswith(("Payment_", "Date_"))}
            os.makedirs(self.output_folder, exist_ok=True)
            page_size = landscape(A4)
            total_rows = max(1, sheet.max_row - 1)
            for row_number in range(2, sheet.max_row + 1):
                if self.stop_flag:
                    self.log.emit("⛔ Stopped."); break
                batch = safe_text(sheet.cell(row_number, headers["BN"]).value).zfill(2)
                sub_div, ref_no = safe_text(sheet.cell(row_number, headers["Sdiv"]).value), safe_text(sheet.cell(row_number, headers["AC No."]).value)
                para_sr = safe_text(sheet.cell(row_number, headers["Para Sr. No."]).value)
                if not para_sr or para_sr.lower() == "nan": continue
                self.log.emit(f"🔍 BN {batch} | SDiv {sub_div} | AC No. {ref_no}")
                temp_pdfs, total_amount, page_number = [], 0.0, 1
                final_pdf = Path(self.output_folder) / f"{para_sr}.pdf"
                # A result belongs only to this run.  Remove any previous
                # result before searching so a no-payment run can never leave
                # an old PDF looking like a newly extracted one.
                if final_pdf.exists():
                    try:
                        final_pdf.unlink()
                        self.log.emit(f"🧹 Removed previous PDF: {final_pdf.name}")
                    except OSError as exc:
                        self.log.emit(f"⚠ Could not remove previous PDF {final_pdf.name}: {exc}")
                        sheet.cell(row_number, status_col).value = "PDF not created - output locked"
                        continue
                resolved_ref = consumer_reference(ref_no, sub_div)
                self.log.emit(f"🎯 Ledger AC used for exact matching: {resolved_ref or '(none)'}")
                for record in by_batch.get(batch, []):
                    try: text = record["path"].read_text(encoding="utf-8", errors="ignore")
                    except OSError as exc: self.log.emit(f"⚠ Could not read {record['path'].name}: {exc}"); continue
                    for page in split_pages(text, ledger_header(text)):
                        if sub_div and not re.search(rf"(?<!\d){re.escape(sub_div)}(?!\d)", page): continue
                        amount, date_dd_mm = extract_payment(page, ref_no, sub_div)
                        if not amount: continue
                        year, total_amount = ledger_year(text, record["year"]), total_amount + float(amount.replace(",", ""))
                        self.log.emit(f"✅ Payment {page_number} | Amount: {amount} | Date: {date_dd_mm}/{year}")
                        pay_index = 1
                        while f"Payment_{pay_index}" in payment_cols and sheet.cell(row_number, payment_cols[f"Payment_{pay_index}"]).value: pay_index += 1
                        pay_header, date_header = f"Payment_{pay_index}", f"Date_{pay_index}"
                        if pay_header not in payment_cols:
                            next_col = sheet.max_column + 1; sheet.cell(1, next_col).value = pay_header; sheet.cell(1, next_col + 1).value = date_header
                            payment_cols[pay_header], payment_cols[date_header] = next_col, next_col + 1
                        sheet.cell(row_number, payment_cols[pay_header]).value = amount
                        sheet.cell(row_number, payment_cols[date_header]).value = f"{date_dd_mm}/{year}" if year else date_dd_mm
                        temp = Path(self.output_folder) / f"temp_{para_sr}_{page_number}.pdf"
                        self.create_pdf(page, temp, f"Para Sr. No. {para_sr}" if page_number == 1 else f"Para Sr. No. {para_sr} ({page_number})", consumer_reference(ref_no, sub_div), amount, date_dd_mm, canvas, page_size, colors)
                        temp_pdfs.append(temp); page_number += 1
                if temp_pdfs:
                    merger = PdfMerger()
                    for temp in temp_pdfs: merger.append(str(temp))
                    merger.write(str(final_pdf)); merger.close()
                    for temp in temp_pdfs:
                        try: temp.unlink()
                        except OSError: pass
                    sheet.cell(row_number, status_col).value, sheet.cell(row_number, total_col).value = "Payment Found & Extracted", total_amount
                    self.log.emit(f"📄 PDF created: {final_pdf.name}")
                else:
                    sheet.cell(row_number, status_col).value, sheet.cell(row_number, total_col).value = "No Payment", 0
                    self.log.emit("❌ No payment found")
                self.progress.emit(int((row_number - 1) / total_rows * 100))
            workbook.save(self.excel_path); self.log.emit("🎯 PROCESS COMPLETED")
        except Exception as exc: self.log.emit(f"ERROR: {exc}")
        finally: self.finished.emit()

    @staticmethod
    def create_pdf(text, output_path, label, ref_no, amount, date_dd_mm, canvas, page_size, colors):
        c = canvas.Canvas(str(output_path), pagesize=page_size); width, height = page_size; lines = text.splitlines()
        if not lines: c.save(); return
        left, top, bottom = 18, 30, 22; usable_width = width - 36; reference_size = 9
        longest = max(c.stringWidth(line, "Courier", reference_size) for line in lines)
        width_size = reference_size * usable_width / max(longest, 1)
        # Fit the complete source page vertically.  The old fixed line gap
        # silently dropped the bottom rows, including a valid matched AC and
        # its payment, from the generated PDF.
        available_height = height - top - bottom
        height_gap = available_height / max(len(lines), 1)
        gap = min(reference_size + 1, height_gap)
        font_size = max(5, min(reference_size, width_size, gap - 1))
        gap = max(font_size + 1, height_gap)
        c.setFont("Courier", font_size); first_width = c.stringWidth(lines[0], "Courier", font_size); c.setFont("Helvetica-Bold", 10); label_width = c.stringWidth(label, "Helvetica-Bold", 10); label_x = left + first_width - label_width; label_y = height - top + gap
        c.drawString(label_x, label_y, label); c.line(label_x, label_y - 1.5, label_x + label_width, label_y - 1.5); c.setFont("Courier", font_size); y = height - top - gap
        for line in lines:
            if ref_no in line and amount in line and date_dd_mm in line:
                c.saveState(); c.setFillAlpha(0.25); c.setFillColor(colors.yellow); c.rect(left - 2, y - 2, c.stringWidth(line, "Courier", font_size) + 4, gap + 2, 0, 1); c.restoreState()
            c.drawString(left, y, line); y -= gap
        c.save()


class PaymentExtract88LPanel(QWidget):
    def __init__(self):
        super().__init__(); self.worker = None
        layout = QVBoxLayout(self); layout.setSpacing(10); layout.addWidget(lbl("88L Payment Extractor", bold=True, color="#7eb8f7")); layout.addWidget(hline())
        group = QGroupBox("Inputs"); form = QVBoxLayout(group)
        self.excel_ed = self.add_path_row(form, "Excel File:", False, "Excel (*.xlsx)")
        template = QPushButton("⬇  Download Template"); template.clicked.connect(self.get_template); form.itemAt(form.count() - 1).layout().addWidget(template)
        self.root_ed = self.add_path_row(form, "88L Folder:", True, "")
        self.output_ed = self.add_path_row(form, "Output Folder:", True, "")
        layout.addWidget(group)
        buttons = QHBoxLayout(); self.start_btn = QPushButton("▶  START"); self.start_btn.setStyleSheet(START_BTN); self.stop_btn = QPushButton("■  STOP"); self.stop_btn.setStyleSheet(STOP_BTN); self.stop_btn.setEnabled(False); self.start_btn.clicked.connect(self.start); self.stop_btn.clicked.connect(self.stop); buttons.addWidget(self.start_btn); buttons.addWidget(self.stop_btn); buttons.addStretch(); layout.addLayout(buttons)
        self.progress = QProgressBar(); layout.addWidget(self.progress); self.log_box = QTextEdit(); self.log_box.setReadOnly(True); layout.addWidget(self.log_box)

    def add_path_row(self, parent, label, folder, file_filter):
        row = QHBoxLayout(); row.addWidget(QLabel(label)); edit = QLineEdit(); edit.setReadOnly(True); edit.setProperty("preferred_width", 245); row.addWidget(edit, 1); button = QPushButton("Browse")
        button.clicked.connect(lambda: self.choose_folder(edit) if folder else self.choose_file(edit, file_filter)); row.addWidget(button); parent.addLayout(row); return edit

    def choose_folder(self, edit):
        path = QFileDialog.getExistingDirectory(self, "Select Folder", edit.text())
        if path: edit.setText(path)

    def choose_file(self, edit, file_filter):
        path, _ = QFileDialog.getOpenFileName(self, "Select Excel File", "", file_filter)
        if path: edit.setText(path)

    def get_template(self):
        download_template(self, headers=["BN", "Sdiv", "AC No.", "Para Sr. No."], default_name="PaymentExtract_88L_Template.xlsx")

    def start(self):
        if not all((self.excel_ed.text(), self.root_ed.text(), self.output_ed.text())):
            QMessageBox.warning(self, "Missing", "Select an Excel file, 88L folder, and output folder first."); return
        self.worker = PaymentExtract88LWorker(self.excel_ed.text(), self.root_ed.text(), self.output_ed.text()); self.worker.log.connect(self.log_box.append); self.worker.progress.connect(self.progress.setValue); self.worker.finished.connect(self.on_finished); self.start_btn.setEnabled(False); self.stop_btn.setEnabled(True); self.worker.start()

    def stop(self):
        if self.worker: self.worker.stop()

    def on_finished(self):
        self.start_btn.setEnabled(True); self.stop_btn.setEnabled(False)
        notify(self, "Completed", "88L payment extraction completed.")
