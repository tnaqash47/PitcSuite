import os
import re
import time
import base64
import threading
import configparser

from PySide6.QtWidgets import (
    QApplication, QWidget, QLabel, QPushButton,
    QLineEdit, QFileDialog, QCheckBox, QTextEdit, QProgressBar,
    QVBoxLayout, QHBoxLayout, QGridLayout, QGroupBox, QMessageBox, QComboBox,
    QSpinBox
)
from PySide6.QtCore import Signal, QObject, QThread

from pitcsuite.ui_helpers import lbl, hline, START_BTN, STOP_BTN, notify
from pitcsuite.templates import download_template
from pitcsuite.config import config_path as _config_path

REPORT_HREF = {
    "meter":      "MeterReadingAndConsumption",
    "billing":    "BillingDetails",
    "adjustment": "BillingAdjustmentDetails",
    "payment":    "PaymentDetails",
    "summary":    "CustomerDataSummary",
}

# PITC billing values can differ from the Excel amount in the final rupee
# digit.  Accept differences of up to 9 rupees when locating and highlighting
# the requested amount.
AMOUNT_DIFFERENCE_TOLERANCE = 9


class PITCWorker(QThread):
    progress = Signal(int)
    log      = Signal(str)

    def __init__(self, excel, report_key, download_dir, profile_dir,
                 restart_after, warmup, print_wait, print_x, print_y,
                 page_from=None, page_to=None, match_mode="first"):
        super().__init__()
        self.excel         = excel
        self.report_key    = report_key
        self.download_dir  = download_dir
        self.profile_dir   = profile_dir
        self.restart_after = restart_after
        self.warmup        = warmup
        self.print_wait    = print_wait
        self.print_x       = print_x
        self.print_y       = print_y
        self.page_from     = page_from
        self.page_to       = page_to
        self.match_mode    = match_mode
        self.running       = True
        self.stopped       = False
        self._driver       = None
        self._driver_lock  = threading.Lock()

    def stop(self):
        self.running = False
        self.stopped = True
        # Closing Chrome from a short-lived helper prevents a Selenium call in
        # the worker from keeping the application UI hostage.  The GUI thread
        # never waits for this operation.
        with self._driver_lock:
            driver = self._driver
        if driver:
            threading.Thread(target=self._close_driver, args=(driver,), daemon=True).start()

    @staticmethod
    def _close_driver(driver):
        try:
            driver.quit()
        except Exception:
            pass

    def run(self):
        try:
            self._run_impl()
        finally:
            with self._driver_lock:
                driver = self._driver
                self._driver = None
            if driver:
                self._close_driver(driver)

    def _run_impl(self):
        try:
            from selenium import webdriver
            from selenium.webdriver.chrome.service import Service
            from selenium.webdriver.chrome.options import Options
            from selenium.webdriver.common.by import By
            from selenium.webdriver.common.keys import Keys
            from selenium.webdriver.support.ui import WebDriverWait
            from selenium.webdriver.support import expected_conditions as EC
            from selenium.common.exceptions import InvalidSessionIdException, StaleElementReferenceException, TimeoutException
            from webdriver_manager.chrome import ChromeDriverManager
            from openpyxl import load_workbook
            from pypdf import PdfReader, PdfWriter
            import pyautogui, fitz, shutil, traceback
            import datetime as _dt
            import pyotp
        except ImportError as e:
            self.log.emit(f"ERROR: Missing library — {e}")
            return

        cfg = configparser.ConfigParser()
        cfg.read(_config_path())
        try:
            USERNAME   = cfg["credentials"]["username"]
            PASSWORD   = cfg["credentials"]["password"]
            OTP_SECRET = cfg["credentials"].get("otp_secret", "")
        except Exception:
            self.log.emit("ERROR: config.ini missing or invalid. Set credentials in ⚙️ Settings first.")
            return

        dl = self.download_dir
        os.makedirs(dl, exist_ok=True)
        os.makedirs(self.profile_dir, exist_ok=True)
        pyautogui.FAILSAFE = False
        pyautogui.PAUSE = 0.3

        # ── OTP ──
        def generate_otp():
            totp = pyotp.TOTP(OTP_SECRET)
            while True:
                remaining = totp.interval - (time.time() % totp.interval)
                if remaining < 10:
                    self.log.emit("Waiting for fresh OTP window…")
                    time.sleep(remaining + 1)
                else:
                    break
            return totp.now()

        # ── PDF helpers ──
        def wait_for_new_pdf(existing, timeout=60):
            start = time.time()
            while self.running and time.time() - start < timeout:
                pdfs = {f for f in os.listdir(dl)
                        if f.lower().endswith(".pdf") and not f.lower().endswith(".crdownload")}
                diff = pdfs - existing
                if diff:
                    path = os.path.join(dl, diff.pop())
                    if os.path.getsize(path) < 5000:
                        raise Exception("Invalid PDF (login/session issue)")
                    return path
                if self.running:
                    time.sleep(0.5)
            if not self.running:
                raise InterruptedError("Download stopped by user")
            raise Exception("PDF not downloaded within timeout")

        def get_unique_name(base_name):
            base_name = str(base_name).strip()
            name = f"{base_name}.pdf"
            if not os.path.exists(os.path.join(dl, name)): return name
            i = 1
            while True:
                name = f"{base_name}_{i}.pdf"
                if not os.path.exists(os.path.join(dl, name)): return name
                i += 1

        def format_month(val):
            if isinstance(val, _dt.datetime):
                return val.strftime("%b - %Y")
            s = str(val).strip()
            for fmt in ("%b-%y", "%b-%Y", "%m-%y", "%m-%Y"):
                try:
                    return _dt.datetime.strptime(s, fmt).strftime("%b - %Y")
                except:
                    pass
            return s

        def amount_as_integer(value):
            """Convert an amount to whole rupees, ignoring decimal places."""
            try:
                from decimal import Decimal, InvalidOperation
                raw = str(value).replace("\u00a0", " ").strip()
                # PDF text may contain spaces inside a formatted number, e.g.
                # ``85, 567.00``. Remove only spaces between digits.
                raw = re.sub(r"(?<=\d)\s+(?=[\d,])", "", raw)
                match = re.search(r"[-+]?\d[\d,]*(?:\.\d+)?", raw)
                if not match:
                    return None
                return int(Decimal(match.group(0).replace(",", "")))
            except (InvalidOperation, ValueError, TypeError):
                return None

        def amount_matches(value, amount, tolerance=AMOUNT_DIFFERENCE_TOLERANCE):
            """Return True when whole-rupee amounts differ by at most 9."""
            value = amount_as_integer(value)
            target = amount_as_integer(amount)
            return value is not None and target is not None and abs(value - target) <= tolerance

        def amount_rects(page, amount):
            target = amount_as_integer(amount)
            if target is None:
                return []
            rects = []
            for word in page.get_text("words"):
                if amount_matches(word[4], target):
                    rects.append(fitz.Rect(word[:4]))
            return rects

        def text_has_amount(text, amount):
            target = amount_as_integer(amount)
            if target is None:
                return False
            # Some PDF extractors put a space after a thousands separator.
            text = re.sub(r"(?<=,)[ \t]+(?=\d)", "", text)
            number_tokens = re.findall(r"(?<![\d.-])[\d,]+(?:\.\d+)?(?![\d.-])", text)
            return any(amount_matches(token, target) for token in number_tokens)

        def debit_amount_in_month_row(page, month, amount):
            """Find the closest debit amount on the row containing the month.

            Billing reports can differ by up to nine rupees from the Excel value. The
            row match deliberately ignores decimal places and prefers the
            closest numeric value, while avoiding amounts from other rows.
            """
            target = amount_as_integer(amount)
            if target is None:
                return []

            month_rects = page.search_for(format_month(month))
            if not month_rects:
                return []
            month_rect = month_rects[0]
            month_y = (month_rect.y0 + month_rect.y1) / 2

            candidates = []
            for word in page.get_text("words"):
                word_rect = fitz.Rect(word[:4])
                word_y = (word_rect.y0 + word_rect.y1) / 2
                if word_rect.intersects(month_rect):
                    continue
                if abs(word_y - month_y) > max(8, month_rect.height * 1.5):
                    continue
                value = amount_as_integer(word[4])
                if value is None or value < 0:
                    continue
                candidates.append((abs(value - target), word_rect))

            if not candidates:
                return []

            candidates.sort(key=lambda item: item[0])
            # A nine-rupee tolerance handles final-digit posting differences
            # while still preferring the closest amount on the matching row.
            return [candidates[0][1]] if candidates[0][0] <= AMOUNT_DIFFERENCE_TOLERANCE else []

        def filter_pdf(pdf_path, month=None, amount=None, page_from=None, page_to=None):
            """Filter by an explicit page range or by month/amount."""
            temp_filtered = pdf_path.replace(".pdf", "_filtered.pdf")

            if page_from is not None and page_to is not None:
                page_doc = fitz.open(pdf_path)
                total = len(page_doc)
                if page_from < 1 or page_to < page_from or page_to > total:
                    page_doc.close()
                    raise Exception(
                        f"Page range {page_from}-{page_to} is outside the downloaded PDF (1-{total})"
                    )
                page_doc.select(list(range(page_from - 1, page_to)))
                page_doc.save(temp_filtered, garbage=4, deflate=True, clean=True)
                page_doc.close()
                os.replace(temp_filtered, pdf_path)
                self.log.emit(f"PDF filtered to pages {page_from}-{page_to}")
                return

            reader = PdfReader(pdf_path)
            fitz_doc = fitz.open(pdf_path)
            writer = PdfWriter()
            month_str = format_month(month).lower() if month else None
            has_month = bool(month_str)
            has_amount = amount_as_integer(amount) is not None
            matching_pages = []
            for page_number, page in enumerate(reader.pages):
                # Use the same extractor as highlighting. Browser-generated
                # PDFs can return incomplete text through pypdf, especially
                # for amounts containing commas and decimal places.
                text = fitz_doc[page_number].get_text("text").lower()
                m = has_month and month_str in text
                a = has_amount and text_has_amount(text, amount)
                # With both filters supplied, require both on the same page.
                # With only an amount supplied, amount matching is sufficient.
                matches = (m and a) if (has_month and has_amount) else (m or a)
                if matches:
                    matching_pages.append(page_number)
            fitz_doc.close()
            if self.match_mode == "second":
                selected_pages = matching_pages[1:2]
            elif self.match_mode == "all":
                selected_pages = matching_pages
            else:
                selected_pages = matching_pages[:1]

            if not selected_pages:
                os.remove(pdf_path)
                if self.match_mode == "second" and matching_pages:
                    raise Exception("Only one matching page found; second match does not exist")
                raise Exception("No matching page found (month/amount)")

            for page_number in selected_pages:
                writer.add_page(reader.pages[page_number])
            with open(temp_filtered, "wb") as f:
                writer.write(f)
            self.log.emit(
                f"Selected {self.match_mode} matching page(s): "
                f"{len(selected_pages)} of {len(matching_pages)} found"
            )
            # highlight
            doc = fitz.open(temp_filtered)
            for page in doc:
                if has_month:
                    month_rects = page.search_for(format_month(month))
                    for rect in month_rects:
                        page.add_highlight_annot(rect).update()
                if has_amount:
                    if has_month:
                        areas = debit_amount_in_month_row(page, month, amount)
                    else:
                        areas = amount_rects(page, amount)
                    for rect in areas:
                        page.add_highlight_annot(rect).update()
            highlighted = pdf_path.replace(".pdf", "_highlighted.pdf")
            doc.save(highlighted, garbage=4, deflate=True, clean=True)
            doc.close()
            os.remove(pdf_path)
            os.remove(temp_filtered)
            shutil.move(highlighted, pdf_path)
            self.log.emit("PDF filtered and highlighted")

        def stamp_sr(pdf_path, sr):
            doc = fitz.open(pdf_path)
            total = len(doc)
            for i, page in enumerate(doc, 1):
                # page.rect is rotation-aware, so the stamp stays at the
                # same visual corner on portrait and landscape pages.
                rect = page.rect
                label = f"Sr No. {sr} ({i}/{total})"
                label_width = fitz.get_text_length(label, fontname="hebo", fontsize=14)
                page.insert_text(
                    (rect.width - label_width - 20, 24),
                    label,
                    fontsize=14,
                    fontname="hebo",
                    color=(0, 0, 0),
                )
            stamped = pdf_path.replace(".pdf", "_stamped.pdf")
            doc.save(stamped, garbage=4, deflate=True, clean=True)
            doc.close()
            os.replace(stamped, pdf_path)

        def start_driver():
            options = Options()
            options.add_argument("--start-maximized")
            options.add_argument("--disable-notifications")
            options.add_argument("--disable-blink-features=AutomationControlled")
            options.add_argument("--disable-dev-shm-usage")
            options.add_argument("--no-sandbox")
            options.add_argument(f"--user-data-dir={self.profile_dir}")
            options.add_experimental_option("prefs", {
                "download.default_directory": dl,
                "download.prompt_for_download": False,
                "download.directory_upgrade": True,
                "safebrowsing.enabled": False,
            })
            drv = webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=options)
            with self._driver_lock:
                self._driver = drv
            return drv, WebDriverWait(drv, 40)

        def do_login(drv, wt):
            drv.get("https://cp.pitc.com.pk")
            wt.until(EC.element_to_be_clickable((By.XPATH, "//button[@data-id='ctl00_ContentPlaceHolder1_cmbSelectDisco']"))).click()
            wt.until(EC.element_to_be_clickable((By.XPATH, "//span[text()='PESCO']"))).click()
            wt.until(EC.presence_of_element_located((By.NAME, "ctl00$ContentPlaceHolder1$txtUsername"))).send_keys(USERNAME)
            wt.until(EC.presence_of_element_located((By.NAME, "ctl00$ContentPlaceHolder1$txtPassword"))).send_keys(PASSWORD)
            wt.until(EC.element_to_be_clickable((By.ID, "btnLogin"))).click()
            if OTP_SECRET:
                self.log.emit("Waiting for OTP…")
                for attempt in range(5):
                    try:
                        otp_box = wt.until(EC.element_to_be_clickable((By.ID, "ctl00_ContentPlaceHolder1_txtOtp")))
                        otp_code = generate_otp()
                        self.log.emit(f"OTP attempt {attempt+1}: {otp_code}")
                        otp_box.send_keys(Keys.CONTROL, "a", Keys.BACKSPACE)
                        otp_box.send_keys(otp_code)
                        wt.until(EC.element_to_be_clickable((By.ID, "ctl00_ContentPlaceHolder1_btnVerify"))).click()
                        WebDriverWait(drv, 10).until(EC.element_to_be_clickable((By.NAME, "ctl00$ContentPlaceHolder1$txtReferenceNumber")))
                        if drv.find_elements(By.NAME, "ctl00$ContentPlaceHolder1$txtReferenceNumber"):
                            self.log.emit("OTP verified ✔")
                            return
                        time.sleep(0.5)
                    except TimeoutException:
                        self.log.emit("OTP not accepted; retrying")
                    except InvalidSessionIdException:
                        raise
                    except Exception as e:
                        self.log.emit(f"OTP retry: {e}")
                    time.sleep(1)
                raise Exception("OTP verification failed")

        def ensure_logged_in(drv, wt):
            if not drv.find_elements(By.NAME, "ctl00$ContentPlaceHolder1$txtReferenceNumber"):
                self.log.emit("Session expired → re-login")
                do_login(drv, wt)

        def login_with_session_recovery(drv, wt):
            try:
                do_login(drv, wt)
                return drv, wt
            except InvalidSessionIdException:
                self.log.emit("Browser session lost during OTP; restarting Chrome")
                try:
                    drv.quit()
                except Exception:
                    pass
                drv, wt = start_driver()
                do_login(drv, wt)
                return drv, wt

        def open_report_link(drv, wt, href_kw):
            """Click the report link with stale-element retry."""
            for attempt in range(5):
                try:
                    link = wt.until(EC.element_to_be_clickable(
                        (By.XPATH, f"//a[contains(@href,'{href_kw}')]")
                    ))
                    drv.execute_script("arguments[0].scrollIntoView({block:'center'});", link)
                    time.sleep(0.5)
                    drv.execute_script("arguments[0].click();", link)
                    return
                except StaleElementReferenceException:
                    self.log.emit(f"Stale element retry {attempt+1}")
                    time.sleep(1)
                except TimeoutException:
                    self.log.emit(f"Timeout retry {attempt+1}")
                    time.sleep(1)
            raise Exception(f"Could not open report link: {href_kw}")

        def missing_ac_status(drv):
            """Return the workbook status when the portal reports no AC data.

            The portal has used several different messages for the same
            condition over time.  Check the visible page text before opening
            a report and again after opening it so a stale/empty report page
            cannot reach the Print button.
            """
            try:
                text = drv.find_element(By.TAG_NAME, "body").text
            except Exception:
                return None

            normalized = re.sub(r"\s+", " ", text).strip().lower()
            missing_phrases = (
                "data not exist",
                "data does not exist",
                "data doesn't exist",
                "no data found",
                "record not found",
                "bill not found",
                "does not belongs",
                "does not belong",
                "wrong ac no",
                "wrong account",
                "invalid account",
            )
            if any(phrase in normalized for phrase in missing_phrases):
                return "data not exist/Wrong AC No."
            return None

        def ac_found_on_page(drv, ac_no):
            """Check that the requested AC number is present in the report."""
            try:
                page = drv.find_element(By.TAG_NAME, "body").text
            except Exception:
                return False

            compact_page = re.sub(r"[^a-z0-9]", "", page.lower())
            compact_ac = re.sub(r"[^a-z0-9]", "", str(ac_no).lower())
            if not compact_ac:
                return False

            candidates = {compact_ac}
            # Excel may expose a numeric AC as 123456.0.  The portal normally
            # renders the underlying integer without the trailing .0.
            if compact_ac.endswith("0") and str(ac_no).strip().endswith(".0"):
                candidates.add(compact_ac[:-1])
            return any(candidate in compact_page for candidate in candidates)

        def wait_for_report_or_missing(drv, href_kw, timeout=8):
            """Wait briefly for either a report link or a missing-AC result."""
            report_xpath = f"//a[contains(@href,'{href_kw}')]"

            def ready(current_driver):
                missing = missing_ac_status(current_driver)
                if missing:
                    return missing
                if current_driver.find_elements(By.XPATH, report_xpath):
                    return "report"
                return False

            try:
                result = WebDriverWait(drv, timeout, poll_frequency=0.25).until(ready)
                return None if result == "report" else result
            except TimeoutException:
                # No report link after the short search window means this AC
                # is unavailable; avoid the old 5 x 40-second retry path.
                return "data not exist/Wrong AC No."

        driver, wait = start_driver()
        driver, wait = login_with_session_recovery(driver, wait)

        wb         = load_workbook(self.excel)
        sheet      = wb.active
        total_rows = sheet.max_row - 1
        processed  = 0
        row        = 2
        href_kw    = REPORT_HREF[self.report_key]

        wait.until(EC.element_to_be_clickable((By.NAME, "ctl00$ContentPlaceHolder1$txtReferenceNumber")))
        if self.warmup > 0:
            time.sleep(self.warmup)

        while True:
            if not self.running: break
            ac = sheet[f"A{row}"].value
            if not ac: break
            sr     = sheet[f"B{row}"].value
            month  = sheet[f"C{row}"].value
            amount = sheet[f"D{row}"].value
            status = sheet[f"E{row}"]

            self.log.emit(f"Row {row} | AC: {ac}")

            try:
                ensure_logged_in(driver, wait)
                report_opened = False

                box = wait.until(EC.element_to_be_clickable((By.NAME, "ctl00$ContentPlaceHolder1$txtReferenceNumber")))
                box.send_keys(Keys.CONTROL, "a", Keys.BACKSPACE)
                box.send_keys(str(ac))

                wait.until(EC.element_to_be_clickable((By.NAME, "ctl00$ContentPlaceHolder1$btnGo"))).click()
                time.sleep(2)

                missing_status = wait_for_report_or_missing(driver, href_kw)
                if missing_status:
                    status.value = missing_status
                    self.log.emit(f"→ {missing_status}; print skipped")
                else:
                    open_report_link(driver, wait, href_kw)
                    report_opened = True
                    try:
                        WebDriverWait(driver, 15).until(
                            lambda drv: missing_ac_status(drv) or ac_found_on_page(drv, ac)
                        )
                    except TimeoutException:
                        pass

                    missing_status = missing_ac_status(driver)
                    if missing_status or not ac_found_on_page(driver, ac):
                        missing_status = missing_status or "data not exist/Wrong AC No."
                        status.value = missing_status
                        self.log.emit(f"→ {missing_status}; AC not found on report page; print skipped")
                    else:
                        wait.until(EC.element_to_be_clickable((By.ID, "ctl00_ContentPlaceHolder1_btnPrint"))).click()

                if missing_status:
                    wb.save(self.excel)
                    try:
                        if report_opened:
                            # The report opens in the current Chrome tab. Do
                            # not use Ctrl+W here: it can close the only
                            # browser window and invalidate the Selenium
                            # session for every remaining Excel row.
                            driver.back()
                            WebDriverWait(driver, 10).until(
                                EC.presence_of_element_located(
                                    (By.NAME, "ctl00$ContentPlaceHolder1$txtReferenceNumber")
                                )
                            )
                    except Exception:
                        pass
                    time.sleep(1)
                else:
                    existing = {f for f in os.listdir(dl) if f.lower().endswith(".pdf")}
                    time.sleep(self.print_wait)

                    pyautogui.click(self.print_x, self.print_y)
                    pyautogui.press(["down"])
                    pyautogui.press(["tab", "tab", "tab", "enter"])

                    new_pdf    = wait_for_new_pdf(existing)
                    base_name  = sr if sr else ac
                    final_name = get_unique_name(base_name)
                    final_path = os.path.join(dl, final_name)
                    os.rename(new_pdf, final_path)

                    if self.page_from is not None and self.page_to is not None:
                        filter_pdf(final_path, page_from=self.page_from, page_to=self.page_to)
                    elif month or amount_as_integer(amount) is not None:
                        filter_pdf(final_path, month, amount)

                    if sr:
                        stamp_sr(final_path, sr)

                    status.value = "Done"
                    self.log.emit(f"✅ PDF created: {final_name}")

                    pyautogui.hotkey("ctrl", "w")
                    time.sleep(2)

            except Exception as e:
                traceback.print_exc()
                status.value = "Error"
                self.log.emit(f"⚠ ERROR: {e}")
                try: pyautogui.hotkey("ctrl", "w")
                except: pass
                time.sleep(3)

            row += 1
            processed += 1
            self.progress.emit(int((processed / max(total_rows, 1)) * 100))

            if processed % self.restart_after == 0:
                self.log.emit("Restarting browser session…")
                driver.quit(); time.sleep(6)
                driver, wait = start_driver()
                driver, wait = login_with_session_recovery(driver, wait)
                wait.until(EC.element_to_be_clickable((By.NAME, "ctl00$ContentPlaceHolder1$txtReferenceNumber")))
                if self.warmup > 0:
                    time.sleep(self.warmup)

        wb.save(self.excel)


class PITCPanel(QWidget):
    REPORT_LABELS = [
        "Meter Reading & Consumption",
        "Billing Detail",
        "Billing Adjustment",
        "Payment Detail",
        "Customer Data Summary",
    ]
    REPORT_KEYS = ["meter", "billing", "adjustment", "payment", "summary"]

    def __init__(self):
        super().__init__()
        self.worker = None
        layout = QVBoxLayout(self)
        layout.setSpacing(10)

        layout.addWidget(lbl("PITC Downloader", bold=True, color="#7eb8f7"))
        layout.addWidget(lbl("OTP login  ·  All report types  ·  Month & Amount filter  ·  Highlight  ·  Para Sr. Stamp", color="#556"))
        layout.addWidget(hline())

        # ── Excel + Template ──
        g1 = QGroupBox("Excel File")
        g1v = QVBoxLayout(g1)
        r_ex = QHBoxLayout()
        self.excel_ed = QLineEdit()
        self.excel_ed.setPlaceholderText("Select Excel file (*.xlsx)...")
        self.excel_ed.setReadOnly(True)
        self.excel_ed.setProperty("preferred_width", 245)
        b1 = QPushButton("Browse"); b1.clicked.connect(self.pick_excel)
        btn_tpl = QPushButton("⬇  Download Template")
        btn_tpl.setStyleSheet("background:#194; border:none; border-radius:4px; padding:5px 10px; color:#fff;")
        btn_tpl.clicked.connect(self.get_template)
        r_ex.addWidget(self.excel_ed); r_ex.addWidget(b1); r_ex.addWidget(btn_tpl)
        g1v.addLayout(r_ex)
        layout.addWidget(g1)

        # ── Report Type ──
        g_rep = QGroupBox("Report Type")
        g_rep_h = QHBoxLayout(g_rep)
        self.report_cb = QComboBox()
        self.report_cb.addItems(self.REPORT_LABELS)
        self.report_cb.setCurrentIndex(4)
        g_rep_h.addWidget(self.report_cb); g_rep_h.addStretch()

        # -- Optional page range --
        g_pages = QGroupBox("Optional Page Range")
        g_pages_h = QGridLayout(g_pages)
        g_pages_h.setHorizontalSpacing(8)
        g_pages_h.setVerticalSpacing(6)
        g_pages_h.addWidget(QLabel("From:"), 0, 0)
        self.page_from_ed = QLineEdit()
        self.page_from_ed.setPlaceholderText("e.g. 2")
        self.page_from_ed.setProperty("preferred_width", 90)
        g_pages_h.addWidget(self.page_from_ed, 0, 1)
        g_pages_h.addWidget(QLabel("To:"), 0, 2)
        self.page_to_ed = QLineEdit()
        self.page_to_ed.setPlaceholderText("e.g. 4")
        self.page_to_ed.setProperty("preferred_width", 90)
        g_pages_h.addWidget(self.page_to_ed, 0, 3)
        g_pages_h.setColumnStretch(1, 1)
        g_pages_h.setColumnStretch(3, 1)
        # -- Match selection --
        g_match = QGroupBox("Matching Pages")
        g_match_h = QHBoxLayout(g_match)
        g_match_h.addWidget(QLabel("Download:"))
        self.match_cb = QComboBox()
        self.match_cb.addItem("First matching page", "first")
        self.match_cb.addItem("Second matching page", "second")
        self.match_cb.addItem("All matching pages", "all")
        self.match_cb.setToolTip("Controls which pages are kept when month/amount matches multiple pages.")
        g_match_h.addWidget(self.match_cb)
        g_match_h.addStretch()
        report_match_row = QHBoxLayout()
        report_match_row.setSpacing(8)
        report_match_row.addWidget(g_rep, 1)
        report_match_row.addWidget(g_match, 1)
        layout.addLayout(report_match_row)

        # Page range and print timing share the same row.
        options_row = QHBoxLayout()
        options_row.setSpacing(8)
        options_row.addWidget(g_pages, 1)
        layout.addLayout(options_row)

        # ── Paths ──
        g2 = QGroupBox("Paths")
        g2v = QHBoxLayout(g2)
        lbl_dl = QLabel("Download Dir:"); lbl_dl.setFixedWidth(82); g2v.addWidget(lbl_dl)
        self.dl_ed = QLineEdit(r"D:\PITC_PDFs"); self.dl_ed.setProperty("preferred_width", 245); g2v.addWidget(self.dl_ed, 1)
        lbl_pr = QLabel("Chrome Profile:"); lbl_pr.setFixedWidth(92); g2v.addWidget(lbl_pr)
        self.pr_ed = QLineEdit(r"C:\ChromeProfiles\PITC"); self.pr_ed.setProperty("preferred_width", 245); g2v.addWidget(self.pr_ed, 1)
        layout.insertWidget(4, g2)

        # ── Timing & Print ──
        g3 = QGroupBox("Timing & Print Dialog")
        g3h = QGridLayout(g3)
        g3h.setHorizontalSpacing(8)
        g3h.setVerticalSpacing(6)

        def spin(val, mn=1, mx=9999):
            s = QSpinBox(); s.setRange(mn, mx); s.setValue(val); return s

        g3h.addWidget(QLabel("Restart"), 0, 0);        self.sp_restart = spin(35, 1, 200); g3h.addWidget(self.sp_restart, 0, 1)
        g3h.addWidget(QLabel("Warmup"), 0, 2);         self.sp_warmup  = spin(0, 0, 200);  g3h.addWidget(self.sp_warmup, 0, 3)
        g3h.addWidget(QLabel("Print Wait (s)"), 1, 0); self.sp_pwait   = spin(6,  1, 200); g3h.addWidget(self.sp_pwait, 1, 1)
        g3h.addWidget(QLabel("Print X"), 2, 0);        self.sp_px      = spin(30);         g3h.addWidget(self.sp_px, 2, 1)
        g3h.addWidget(QLabel("Print Y"), 2, 2);        self.sp_py      = spin(165);        g3h.addWidget(self.sp_py, 2, 3)
        g3h.setColumnStretch(1, 1)
        g3h.setColumnStretch(3, 1)
        options_row.addWidget(g3, 1)

        # ── Buttons ──
        bh = QHBoxLayout()
        self.btn_start = QPushButton("▶  START"); self.btn_start.setStyleSheet(START_BTN)
        self.btn_stop  = QPushButton("■  STOP");  self.btn_stop.setStyleSheet(STOP_BTN); self.btn_stop.setEnabled(False)
        self.btn_start.clicked.connect(self.run); self.btn_stop.clicked.connect(self.halt)
        bh.addWidget(self.btn_start); bh.addWidget(self.btn_stop); bh.addStretch()
        layout.addLayout(bh)

        self.progress = QProgressBar(); layout.addWidget(self.progress)
        self.log_box = QTextEdit(); self.log_box.setReadOnly(True)
        layout.addWidget(self.log_box)

    def pick_excel(self):
        f, _ = QFileDialog.getOpenFileName(self, "Excel File", "", "Excel (*.xlsx)")
        if f: self.excel_ed.setText(f)

    def get_template(self):
        download_template(
            self,
            headers      = ["AC No", "Sr No", "Month", "Amount", "Status"],
            default_name = "PITC_Downloader_Template.xlsx"
        )

    def run(self):
        excel_path = self.excel_ed.text().strip()
        if not excel_path:
            QMessageBox.warning(self, "Excel File Required", "Please select an Excel file before starting.")
            return
        if not os.path.isfile(excel_path):
            QMessageBox.warning(self, "File Not Found", f"The selected Excel file does not exist:\n{excel_path}")
            return

        page_from_text = self.page_from_ed.text().strip()
        page_to_text = self.page_to_ed.text().strip()
        if bool(page_from_text) != bool(page_to_text):
            QMessageBox.warning(self, "Invalid Page Range", "Enter both Pages from and Pages to, or leave both blank.")
            return

        page_from = page_to = None
        if page_from_text:
            try:
                page_from = int(page_from_text)
                page_to = int(page_to_text)
            except ValueError:
                QMessageBox.warning(self, "Invalid Page Range", "Page numbers must be whole numbers.")
                return
            if page_from < 1 or page_to < page_from:
                QMessageBox.warning(self, "Invalid Page Range", "Pages to must be greater than or equal to Pages from.")
                return

        self.btn_start.setEnabled(False); self.btn_stop.setEnabled(True)
        self.log_box.clear()
        self.worker = PITCWorker(
            excel         = self.excel_ed.text(),
            report_key    = self.REPORT_KEYS[self.report_cb.currentIndex()],
            download_dir  = self.dl_ed.text(),
            profile_dir   = self.pr_ed.text(),
            restart_after = self.sp_restart.value(),
            warmup        = self.sp_warmup.value(),
            print_wait    = self.sp_pwait.value(),
            print_x       = self.sp_px.value(),
            print_y       = self.sp_py.value(),
            page_from     = page_from,
            page_to       = page_to,
            match_mode    = self.match_cb.currentData(),
        )
        self.worker.log.connect(self.log_box.append)
        self.worker.progress.connect(self.progress.setValue)
        self.worker.finished.connect(self.done)
        self.worker.start()

    def halt(self):
        if self.worker and self.worker.isRunning():
            self.worker.stop()
            self.btn_stop.setEnabled(False)
            self.log_box.append("⛔ Stop requested; closing Chrome safely…")

    def done(self):
        self.btn_start.setEnabled(True); self.btn_stop.setEnabled(False)
        if not self.worker or not self.worker.stopped:
            notify(self, "Done", "PITC Downloader completed.")


# ──────────────────────────────────────────────────────────────
#  TOOL 2 — E/Bills Downloader (BillSoft.py)
# ──────────────────────────────────────────────────────────────
