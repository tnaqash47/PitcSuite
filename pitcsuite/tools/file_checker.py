import os
import re
import time
import base64
import threading
import configparser

from PySide6.QtWidgets import (
    QApplication, QWidget, QLabel, QPushButton,
    QLineEdit, QFileDialog, QCheckBox, QTextEdit, QProgressBar,
    QVBoxLayout, QHBoxLayout, QGroupBox, QMessageBox, QComboBox,
    QSpinBox
)
from PySide6.QtCore import Signal, QObject, QThread

from pitcsuite.ui_helpers import lbl, hline, START_BTN, STOP_BTN, notify
from pitcsuite.templates import download_template
from pitcsuite.config import config_path as _config_path

class SrNoCheckerWorker(QThread):
    progress = Signal(str)
    done = Signal(bool, str)
    failed = Signal(str)

    def __init__(self, excel_path, folder_path):
        super().__init__()
        self.excel_path = excel_path
        self.folder_path = folder_path
        self.stop_requested = False

    def stop(self):
        self.stop_requested = True

    def run(self):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(self.excel_path); ws = wb.active
            files_in_folder = os.listdir(self.folder_path)
            row = 2
            while not self.stop_requested:
                sr_no = ws[f"A{row}"].value
                if sr_no is None: break
                found = any(os.path.splitext(f)[0] == str(sr_no) for f in files_in_folder)
                ws[f"B{row}"] = "Found" if found else "Not Found"
                self.progress.emit(f"Checking Sr. No.: {sr_no}")
                row += 1
            wb.save(self.excel_path)
            self.done.emit(self.stop_requested, "Process stopped." if self.stop_requested else "Process completed successfully.")
        except Exception as e:
            self.failed.emit(str(e))


class SrNoCheckerPanel(QWidget):
    def __init__(self):
        super().__init__()
        self.worker = None
        layout = QVBoxLayout(self); layout.setSpacing(10)
        layout.addWidget(lbl("File Existence Checker", bold=True, color="#7eb8f7"))
        layout.addWidget(lbl("Checks if Sr. No. from Excel column A exists as a file in the selected folder.", color="#556"))
        layout.addWidget(hline())

        g = QGroupBox("Inputs")
        gv = QVBoxLayout(g)

        label_w = 70
        r1 = QHBoxLayout(); lbl_excel = QLabel("Excel File:"); lbl_excel.setFixedWidth(label_w); r1.addWidget(lbl_excel)
        self.excel_ed = QLineEdit(); self.excel_ed.setReadOnly(True); self.excel_ed.setProperty("preferred_width", 245); r1.addWidget(self.excel_ed, 1)
        b1 = QPushButton("Browse"); b1.clicked.connect(self._browse_excel); r1.addWidget(b1)
        gv.addLayout(r1)

        r2 = QHBoxLayout(); lbl_folder = QLabel("Folder:"); lbl_folder.setFixedWidth(label_w); r2.addWidget(lbl_folder)
        self.folder_ed = QLineEdit(); self.folder_ed.setReadOnly(True); self.folder_ed.setProperty("preferred_width", 245); r2.addWidget(self.folder_ed, 1)
        b2 = QPushButton("Browse"); b2.clicked.connect(self._browse_folder); r2.addWidget(b2)
        gv.addLayout(r2)
        layout.addWidget(g)

        bh = QHBoxLayout()
        self.btn_start = QPushButton("▶  START"); self.btn_start.setStyleSheet(START_BTN)
        self.btn_stop  = QPushButton("■  STOP");  self.btn_stop.setStyleSheet(STOP_BTN)
        self.btn_start.clicked.connect(self._start); self.btn_stop.clicked.connect(self._stop)
        bh.addWidget(self.btn_start); bh.addWidget(self.btn_stop); bh.addStretch()
        layout.addLayout(bh)

        self.status_lbl = QLabel("Status: Idle"); self.status_lbl.setStyleSheet("color:#4a90d9;")
        layout.addWidget(self.status_lbl)

    def _browse_excel(self):
        f, _ = QFileDialog.getOpenFileName(self, "Excel File", "", "Excel Files (*.xlsx)")
        if f: self.excel_ed.setText(f)

    def _browse_folder(self):
        p = QFileDialog.getExistingDirectory(self, "Select Folder")
        if p: self.folder_ed.setText(p)

    def _start(self):
        if not self.excel_ed.text() or not self.folder_ed.text():
            notify(self, "Error", "Please select Excel file and Folder", critical=True); return
        self.btn_start.setEnabled(False); self.btn_stop.setEnabled(True)
        self.worker = SrNoCheckerWorker(self.excel_ed.text(), self.folder_ed.text())
        self.worker.progress.connect(self.status_lbl.setText)
        self.worker.done.connect(self._done); self.worker.failed.connect(self._failed)
        self.worker.finished.connect(lambda: (self.btn_start.setEnabled(True), self.btn_stop.setEnabled(False)))
        self.worker.start()

    def _stop(self):
        if self.worker and self.worker.isRunning():
            self.worker.stop(); self.status_lbl.setText("Status: Stopping…")

    def _done(self, stopped, message):
        self.status_lbl.setText("Status: Stopped" if stopped else "Status: Completed")
        if not stopped: notify(self, "Done", message)

    def _failed(self, message):
        self.status_lbl.setText("Status: Failed"); notify(self, "Error", message, critical=True)


# ──────────────────────────────────────────────────────────────
#  TOOL 11 — Double CR Checker (DoubleCRchecker.py)
# ──────────────────────────────────────────────────────────────
