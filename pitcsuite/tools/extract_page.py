import os
import re
import time
import base64
import threading
import configparser

from PySide6.QtWidgets import (
    QApplication, QWidget, QLabel, QPushButton,
    QLineEdit, QFileDialog, QCheckBox, QTextEdit, QProgressBar,
    QVBoxLayout, QHBoxLayout, QGroupBox, QMessageBox, QComboBox,
    QSpinBox
)
from PySide6.QtCore import Signal, QObject, QThread

from pitcsuite.ui_helpers import lbl, hline, START_BTN, STOP_BTN
from pitcsuite.templates import download_template
from pitcsuite.config import config_path as _config_path

class ExtractPageWorkerSignals(QObject):
    log      = Signal(str)
    progress = Signal(int, int)
    finished = Signal()


class ExtractPagePanel(QWidget):
    def __init__(self):
        super().__init__()
        self.signals = ExtractPageWorkerSignals()
        self.signals.log.connect(lambda m: self.log_box.append(m))
        self.signals.progress.connect(self.update_progress)
        self.signals.finished.connect(self.on_finished)

        layout = QVBoxLayout(self); layout.setSpacing(10)
        layout.addWidget(lbl("Extract Specific PITC Page", bold=True, color="#7eb8f7"))
        layout.addWidget(lbl("Search Criteria: month e.g. Feb-26  or  any number  ·  Extracts matching page  ·  Highlights  ·  Stamps Para Sr. No.", color="#556"))
        layout.addWidget(hline())

        g = QGroupBox("Inputs")
        gv = QVBoxLayout(g)

        r1 = QHBoxLayout(); r1.addWidget(QLabel("PDF Folder:"))
        self.pdf_ed = QLineEdit(); self.pdf_ed.setReadOnly(True); r1.addWidget(self.pdf_ed, 1)
        b1 = QPushButton("Browse"); b1.clicked.connect(lambda: (p := QFileDialog.getExistingDirectory(self,"PDF Folder")) and self.pdf_ed.setText(p)); r1.addWidget(b1)
        gv.addLayout(r1)

        r2 = QHBoxLayout(); r2.addWidget(QLabel("Excel File:"))
        self.excel_ed = QLineEdit(); self.excel_ed.setReadOnly(True); r2.addWidget(self.excel_ed, 1)
        b2 = QPushButton("Browse"); b2.clicked.connect(lambda: (p := QFileDialog.getOpenFileName(self,"Excel","","Excel (*.xlsx)")[0]) and self.excel_ed.setText(p)); r2.addWidget(b2)
        btn_tpl = QPushButton("⬇  Download Template")
        btn_tpl.setStyleSheet("background:#194; border:none; border-radius:4px; padding:5px 10px; color:#fff;")
        btn_tpl.clicked.connect(self.get_template); r2.addWidget(btn_tpl)
        gv.addLayout(r2)

        r3 = QHBoxLayout(); r3.addWidget(QLabel("Output Folder:"))
        self.out_ed = QLineEdit(r"D:\ExtractedPages"); r3.addWidget(self.out_ed, 1)
        b3 = QPushButton("Change"); b3.clicked.connect(lambda: (p := QFileDialog.getExistingDirectory(self,"Output")) and self.out_ed.setText(p)); r3.addWidget(b3)
        gv.addLayout(r3)

        layout.addWidget(g)

        bh = QHBoxLayout()
        self.btn_start = QPushButton("▶  START"); self.btn_start.setStyleSheet(START_BTN)
        self.btn_stop  = QPushButton("■  STOP");  self.btn_stop.setStyleSheet(STOP_BTN); self.btn_stop.setEnabled(False)
        self.btn_start.clicked.connect(self.start); self.btn_stop.clicked.connect(self.stop)
        bh.addWidget(self.btn_start); bh.addWidget(self.btn_stop); bh.addStretch()
        layout.addLayout(bh)

        ph = QHBoxLayout()
        self.lbl_prog = QLabel("0 / 0"); ph.addWidget(self.lbl_prog); ph.addStretch()
        layout.addLayout(ph)
        self.progress = QProgressBar(); layout.addWidget(self.progress)
        self.log_box = QTextEdit(); self.log_box.setReadOnly(True); layout.addWidget(self.log_box)

        self._stop_flag = False

    def get_template(self):
        download_template(
            self,
            headers      = ["AC No", "Search Criteria", "Para Sr. No."],
            default_name = "ExtractPage_Template.xlsx"
        )

    def update_progress(self, cur, total):
        self.progress.setMaximum(total); self.progress.setValue(cur)
        self.lbl_prog.setText(f"{cur} / {total}")

    def on_finished(self):
        self.btn_start.setEnabled(True); self.btn_stop.setEnabled(False)
        QMessageBox.information(self, "Done", "Extraction completed.")

    def start(self):
        if not self.pdf_ed.text() or not self.excel_ed.text():
            QMessageBox.warning(self, "Missing", "Select PDF folder and Excel file."); return
        self._stop_flag = False
        self.btn_start.setEnabled(False); self.btn_stop.setEnabled(True)
        self.log_box.clear()
        threading.Thread(target=self.process, daemon=True).start()

    def stop(self):
        self._stop_flag = True; self.log_box.append("⛔ Stopping…")

    def process(self):
        try:
            import fitz
            from openpyxl import load_workbook
            import datetime as _dt
        except ImportError as e:
            self.signals.log.emit(f"ERROR: Missing library — {e}")
            self.signals.finished.emit(); return

        def parse_criteria(val):
            """Feb-26 → 'Feb - 2026'.  Pure number or anything else → str as-is."""
            s = str(val).strip()
            for fmt in ("%b-%y", "%b-%Y", "%m-%y", "%m-%Y"):
                try:
                    return _dt.datetime.strptime(s, fmt).strftime("%b - %Y")
                except:
                    pass
            return s

        pdf_folder = self.pdf_ed.text()
        out_folder = self.out_ed.text()
        os.makedirs(out_folder, exist_ok=True)

        wb = load_workbook(self.excel_ed.text()); ws = wb.active
        last_row = ws.max_row
        self.signals.progress.emit(0, last_row - 1)

        for row in range(2, last_row + 1):
            if self._stop_flag: break

            ac_no     = str(ws[f"A{row}"].value or "").strip()
            criteria  = str(ws[f"B{row}"].value or "").strip()
            para_text = str(ws[f"C{row}"].value or "").strip()

            if not ac_no: continue

            search_text = parse_criteria(criteria)
            self.signals.log.emit(f"AC: {ac_no}  |  Searching: '{search_text}'")

            filename = ac_no if ac_no.lower().endswith(".pdf") else ac_no + ".pdf"
            pdf_path = os.path.join(pdf_folder, filename)

            if not os.path.exists(pdf_path):
                ws[f"D{row}"] = "File Not Found"
                self.signals.log.emit("❌ File Not Found")
                self.signals.progress.emit(row - 1, last_row - 1)
                continue

            self.signals.log.emit("✔ File Found")

            try:
                doc   = fitz.open(pdf_path)
                found = False

                for pg_idx, page in enumerate(doc):
                    areas = page.search_for(search_text)
                    # fallback: try compact form e.g. "Feb-2026" if spaced form not found
                    if not areas and " - " in search_text:
                        areas = page.search_for(search_text.replace(" - ", "-"))

                    if areas:
                        found = True
                        self.signals.log.emit(f"✔ Match on page {pg_idx + 1}")

                        # highlight all matches
                        for rect in areas:
                            page.add_highlight_annot(rect).update()

                        # stamp Para Sr. No. top-right
                        if para_text and para_text.lower() != "none":
                            stamp = f"Para Sr. No. {para_text}"
                            pw    = page.rect.width
                            tw    = fitz.get_text_length(stamp, fontname="helv", fontsize=14)
                            page.insert_text(
                                fitz.Point(pw - tw - 30, 26),
                                stamp,
                                fontsize=14,
                                fontname="helv",
                                color=(0, 0, 0),
                            )

                        # save single extracted page
                        out_doc = fitz.open()
                        out_doc.insert_pdf(doc, from_page=pg_idx, to_page=pg_idx)

                        out_path = os.path.join(out_folder, filename)
                        if os.path.exists(out_path):
                            base, ext = os.path.splitext(filename)
                            i = 1
                            while os.path.exists(os.path.join(out_folder, f"{base}_{i}{ext}")):
                                i += 1
                            out_path = os.path.join(out_folder, f"{base}_{i}{ext}")

                        out_doc.save(out_path, garbage=4, deflate=True)
                        out_doc.close()

                        ws[f"D{row}"] = f"Extracted page {pg_idx + 1}"
                        self.signals.log.emit("✅ Extracted & highlighted")
                        break

                doc.close()

                if not found:
                    ws[f"D{row}"] = "Not Found"
                    self.signals.log.emit(f"❌ '{search_text}' not found in PDF")

            except Exception as e:
                ws[f"D{row}"] = "Error"
                self.signals.log.emit(f"⚠ {e}")

            self.signals.progress.emit(row - 1, last_row - 1)

        wb.save(self.excel_ed.text())
        self.signals.finished.emit()

# ──────────────────────────────────────────────────────────────
#  MAIN WINDOW
# ──────────────────────────────────────────────────────────────
