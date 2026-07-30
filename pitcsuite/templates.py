from PySide6.QtWidgets import QFileDialog
from pitcsuite.ui_helpers import notify

def download_template(parent, headers, default_name, sample_rows=None):
    from openpyxl import Workbook as _WB
    from openpyxl.styles import Font, PatternFill, Alignment
    path, _ = QFileDialog.getSaveFileName(
        parent, "Save Template As", default_name, "Excel (*.xlsx)"
    )
    if not path:
        return
    wb = _WB(); ws = wb.active; ws.title = "Data"
    hdr_font  = Font(bold=True, color="FFFFFF")
    hdr_fill  = PatternFill("solid", fgColor="1F4E79")
    hdr_align = Alignment(horizontal="center")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = hdr_font; cell.fill = hdr_fill; cell.alignment = hdr_align
        ws.column_dimensions[cell.column_letter].width = max(len(h) + 6, 16)
    wb.save(path)
    notify(parent, "Template Saved", f"Template saved to:\n{path}")


# ──────────────────────────────────────────────────────────────
#  SETTINGS PANEL — edit & save config.ini from inside the app
# ──────────────────────────────────────────────────────────────
